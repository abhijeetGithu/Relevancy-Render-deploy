import csv
import io
import re
from typing import Dict, List


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip().lower()


def _normalize_url(value: str) -> str:
    return (value or "").strip().lower()


def _is_excel_ooxml(filename: str) -> bool:
    n = (filename or "").lower()
    return n.endswith(".xlsx") or n.endswith(".xlsm")


def parse_ground_truth(filename: str, content: bytes, queries_only: bool = False) -> List[Dict]:
    if _is_excel_ooxml(filename):
        return _parse_xlsx_ground_truth(content, queries_only)
    if filename.lower().endswith(".txt"):
         return _parse_txt_queries(content)
    return _parse_csv_ground_truth(content.decode("utf-8"), queries_only)


def _parse_txt_queries(content: bytes) -> List[Dict]:
    text = content.decode("utf-8")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ValueError("Query file is empty.")
    return [{"query": line, "expected_title": "", "expected_url": ""} for line in lines]


def _parse_csv_ground_truth(content: str, queries_only: bool) -> List[Dict]:
    reader = csv.DictReader(io.StringIO(content))
    if not reader.fieldnames:
        # Fallback for headerless CSV if it's just one column? 
        # But standard CSV usually has headers. If queries_only, maybe just first column?
        # Let's stick to requiring "Query" header for CSV to be safe/consistent.
        raise ValueError("CSV is empty.")

    headers = [h.strip() for h in reader.fieldnames]
    header_map = {h.lower(): h for h in headers}
    query_key = header_map.get("query")
    
    if not query_key:
         # Attempt to guess if single column
        if len(headers) == 1 and queries_only:
             query_key = headers[0]
        else:
            raise ValueError("CSV must have a 'Query' column.")

    title_key = header_map.get("expected title")
    url_key = header_map.get("expected url")

    if not queries_only:
        if not title_key or not url_key:
            raise ValueError("CSV headers must be: Query, Expected Title, Expected URL.")

    parsed: List[Dict] = []
    title_present = []
    url_present = []

    for row in reader:
        query = (row.get(query_key) or "").strip()
        title = (row.get(title_key) or "").strip() if title_key else ""
        url = (row.get(url_key) or "").strip() if url_key else ""
        if not query:
            continue
        parsed.append({"query": query, "expected_title": title, "expected_url": url})
        title_present.append(bool(title))
        url_present.append(bool(url))

    if not parsed:
        raise ValueError("No valid queries found in file.")

    if not queries_only:
        if any(title_present) and not all(title_present):
            raise ValueError("Expected Title must be present for all rows if provided.")
        if any(url_present) and not all(url_present):
            raise ValueError("Expected URL must be present for all rows if provided.")
        if not any(title_present) and not any(url_present):
            raise ValueError("Provide Expected Title for all rows or Expected URL for all rows.")

    return parsed


def _parse_xlsx_ground_truth(content: bytes, queries_only: bool) -> List[Dict]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:
        raise ValueError("XLSX support requires openpyxl. Install with: pip install openpyxl") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("XLSX is empty.")

    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    header_map = {h.lower(): idx for idx, h in enumerate(header)}
    
    if "query" not in header_map:
         if len(header) == 1 and queries_only:
             # Assume single column is query
             header_map["query"] = 0
         else:
            raise ValueError("XLSX must have a 'Query' column.")

    if not queries_only:
        if "expected title" not in header_map or "expected url" not in header_map:
            raise ValueError("XLSX headers must be: Query, Expected Title, Expected URL.")

    parsed: List[Dict] = []
    title_present = []
    url_present = []

    for row in rows[1:]:
        if not row:
            continue
        
        # Safe access to columns
        q_idx = header_map.get("query")
        t_idx = header_map.get("expected title")
        u_idx = header_map.get("expected url")

        query = str(row[q_idx]).strip() if q_idx is not None and len(row) > q_idx and row[q_idx] is not None else ""
        title = str(row[t_idx]).strip() if t_idx is not None and len(row) > t_idx and row[t_idx] is not None else ""
        url = str(row[u_idx]).strip() if u_idx is not None and len(row) > u_idx and row[u_idx] is not None else ""

        if not query:
            continue
        parsed.append({"query": query, "expected_title": title, "expected_url": url})
        title_present.append(bool(title))
        url_present.append(bool(url))

    if not parsed:
        raise ValueError("No valid queries found in file.")

    if not queries_only:
        if any(title_present) and not all(title_present):
            raise ValueError("Expected Title must be present for all rows if provided.")
        if any(url_present) and not all(url_present):
            raise ValueError("Expected URL must be present for all rows if provided.")
        if not any(title_present) and not any(url_present):
            raise ValueError("Provide Expected Title for all rows or Expected URL for all rows.")

    return parsed


def find_rank(results: List[Dict], expected_title: str, expected_url: str) -> int:
    expected_norm = _normalize_text(expected_title)
    expected_url_norm = _normalize_url(expected_url)
    for entry in results:
        title = _normalize_text(entry.get("title", ""))
        url = _normalize_url(entry.get("url", ""))
        
        # Check URL match first (if expected URL is provided)
        if expected_url_norm and expected_url_norm in url:
            return entry.get("rank", -1)
            
        # Check Title match (if expected Title is provided)
        # Ensure title is not empty to avoid false positives with empty strings
        if expected_norm and title:
            if expected_norm in title or title in expected_norm:
                return entry.get("rank", -1)
                
    return -1


def parse_dqe_sheet(filename: str, content: bytes) -> List[Dict]:
    """Parse a DQE search evaluation output sheet.

    The sheet has multiple rows per query (one per SU result).
    ``Generated_Query``, ``Original_Title``, ``Original_URL`` appear only on
    the first row of each query group; subsequent rows leave them blank.

    Returns a list of dicts, one per unique query::

        {
            "query": str,
            "expected_title": str,
            "expected_url": str,
            "su_results": [{"title": str, "url": str, "rank": int}, ...],
        }
    """
    if _is_excel_ooxml(filename):
        rows = _read_xlsx_rows(content)
    else:
        rows = _read_csv_rows(content.decode("utf-8"))

    if not rows:
        raise ValueError("DQE sheet is empty.")

    entries: List[Dict] = []
    current: Dict | None = None

    for row in rows:
        query = (row.get("generated_query") or "").strip()
        if query:
            if current is not None:
                entries.append(current)
            current = {
                "query": query,
                "expected_title": (row.get("original_title") or "").strip(),
                "expected_url": (row.get("original_url") or "").strip(),
                "su_results": [],
            }

        if current is None:
            continue

        result_title = (row.get("result_title") or "").strip()
        result_url = (row.get("result_url") or "").strip()
        result_rank_raw = (row.get("result_rank") or "").strip()
        try:
            result_rank = int(float(result_rank_raw)) if result_rank_raw else len(current["su_results"]) + 1
        except (ValueError, TypeError):
            result_rank = len(current["su_results"]) + 1

        if (result_title or result_url) and len(current["su_results"]) < 10:
            current["su_results"].append({
                "title": result_title,
                "url": result_url,
                "rank": result_rank,
            })

    if current is not None:
        entries.append(current)

    if not entries:
        raise ValueError("No valid queries found in DQE sheet.")

    return entries


def _read_csv_rows(content: str) -> List[Dict]:
    reader = csv.DictReader(io.StringIO(content))
    if not reader.fieldnames:
        return []
    norm_map = {h: h.strip().lower().replace(" ", "_") for h in reader.fieldnames}
    result = []
    for row in reader:
        result.append({norm_map[k]: (v or "").strip() for k, v in row.items()})
    return result


def _read_xlsx_rows(content: bytes) -> List[Dict]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:
        raise ValueError("XLSX support requires openpyxl.") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    # Try "Search Results" sheet first (DQE XLSX format), fallback to active sheet
    sheet = workbook["Search Results"] if "Search Results" in workbook.sheetnames else workbook.active
    raw_rows = list(sheet.iter_rows(values_only=True))
    if not raw_rows:
        return []

    headers = [
        str(cell).strip().lower().replace(" ", "_") if cell is not None else ""
        for cell in raw_rows[0]
    ]
    result = []
    for row in raw_rows[1:]:
        if not row:
            continue
        entry = {}
        for idx, h in enumerate(headers):
            val = str(row[idx]).strip() if idx < len(row) and row[idx] is not None else ""
            entry[h] = val
        result.append(entry)
    return result


def validate_dqe_columns(filename: str, content: bytes) -> None:
    """Validate that a file has the expected DQE sheet columns."""
    if _is_excel_ooxml(filename):
        rows = _read_xlsx_rows(content)
    else:
        rows = _read_csv_rows(content.decode("utf-8"))

    if not rows:
        raise ValueError("File is empty.")

    keys = set(rows[0].keys())
    required = {"generated_query", "result_title", "result_url"}
    missing = required - keys
    if missing:
        raise ValueError(
            f"Missing required columns: {', '.join(sorted(missing))}. "
            f"Expected DQE search evaluation sheet with columns: "
            f"Generated_Query, Original_Title, Original_URL, Result_Title, Result_URL, Result_Rank"
        )


def _get_original_csv_headers(content: str) -> List[str]:
    """Get original (non-normalized) headers from CSV content."""
    reader = csv.DictReader(io.StringIO(content))
    return [h.strip() for h in reader.fieldnames] if reader.fieldnames else []


def _get_original_xlsx_headers(content: bytes) -> List[str]:
    """Get original (non-normalized) headers from XLSX content."""
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:
        raise ValueError("XLSX support requires openpyxl.") from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook["Search Results"] if "Search Results" in workbook.sheetnames else workbook.active
    raw_rows = list(sheet.iter_rows(values_only=True, max_row=1))
    if not raw_rows:
        return []
    return [str(cell).strip() if cell is not None else "" for cell in raw_rows[0]]


def extract_file_columns(filename: str, content: bytes) -> Dict:
    """Extract column names and preview rows for RS column mapping.

    Returns dict with columns (original + normalized names), preview rows, total_rows.
    """
    if _is_excel_ooxml(filename):
        original_headers = _get_original_xlsx_headers(content)
        rows = _read_xlsx_rows(content)
    else:
        text = content.decode("utf-8")
        original_headers = _get_original_csv_headers(text)
        rows = _read_csv_rows(text)

    if not rows:
        raise ValueError("File is empty.")

    normalized_headers = list(rows[0].keys())

    columns = []
    for i, orig in enumerate(original_headers):
        norm = normalized_headers[i] if i < len(normalized_headers) else orig.strip().lower().replace(" ", "_")
        columns.append({"original": orig, "normalized": norm})

    return {
        "columns": columns,
        "preview": rows[:5],
        "total_rows": len(rows),
    }


def parse_rs_sheet(
    filename: str,
    content: bytes,
    column_mapping: Dict[str, str],
    has_ground_truth: bool = True,
) -> List[Dict]:
    """Parse an RS evaluation sheet with user-specified column mapping.

    column_mapping keys (values are normalized column names):
        query_col, result_title_col, result_url_col,
        result_rank_col, expected_title_col, expected_url_col

    Returns same format as parse_dqe_sheet.
    """
    if _is_excel_ooxml(filename):
        rows = _read_xlsx_rows(content)
    else:
        rows = _read_csv_rows(content.decode("utf-8"))

    if not rows:
        raise ValueError("File is empty.")

    query_col = column_mapping.get("query_col", "")
    result_title_col = column_mapping.get("result_title_col", "")
    result_url_col = column_mapping.get("result_url_col", "")
    result_rank_col = column_mapping.get("result_rank_col", "")
    expected_title_col = column_mapping.get("expected_title_col", "") if has_ground_truth else ""
    expected_url_col = column_mapping.get("expected_url_col", "") if has_ground_truth else ""

    if not query_col:
        raise ValueError("Query column mapping is required.")
    if not result_title_col and not result_url_col:
        raise ValueError("At least Result Title or Result URL column is required.")

    entries: List[Dict] = []
    current: Dict | None = None

    for row in rows:
        query = (row.get(query_col) or "").strip()
        if query:
            if current is not None:
                entries.append(current)
            current = {
                "query": query,
                "expected_title": (row.get(expected_title_col) or "").strip() if expected_title_col else "",
                "expected_url": (row.get(expected_url_col) or "").strip() if expected_url_col else "",
                "su_results": [],
            }

        if current is None:
            continue

        result_title = (row.get(result_title_col) or "").strip() if result_title_col else ""
        result_url = (row.get(result_url_col) or "").strip() if result_url_col else ""
        result_rank_raw = (row.get(result_rank_col) or "").strip() if result_rank_col else ""
        try:
            result_rank = int(float(result_rank_raw)) if result_rank_raw else len(current["su_results"]) + 1
        except (ValueError, TypeError):
            result_rank = len(current["su_results"]) + 1

        if (result_title or result_url) and len(current["su_results"]) < 10:
            current["su_results"].append({
                "title": result_title,
                "url": result_url,
                "rank": result_rank,
            })

    if current is not None:
        entries.append(current)

    if not entries:
        raise ValueError("No valid queries found in file.")

    return entries


def parse_google_results_from_content(content: str) -> Dict[str, List[Dict]]:
    """
    Parse Google results from CSV content string.
    
    Expected CSV format:
    query,expected_title,expected_url,google_result_title,google_result_url
    
    Multiple rows per query (up to 10 results). Empty query cells indicate continuation rows.
    
    Returns:
        Dict[str, List[Dict]]: Mapping of query -> list of result dicts with 'title' and 'url' keys
    """
    import logging
    logger = logging.getLogger(__name__)
    
    google_results_map: Dict[str, List[Dict]] = {}
    current_query = None
    
    try:
        reader = csv.DictReader(io.StringIO(content))
        
        if not reader.fieldnames:
            logger.warning("CSV has no headers")
            return {}
        
        # Check if required columns exist
        # We need to be flexible with headers (case insensitive, whitespace)
        headers = [h.strip().lower() for h in reader.fieldnames] if reader.fieldnames else []
        header_map = {h: i for i, h in enumerate(headers)}
        
        logger.info(f"CSV headers (lowercase): {headers}")
        
        google_title_idx = -1
        google_url_idx = -1
        
        # Try to find columns
        possible_title_headers = ["google_result_title", "google result title", "google title"]
        possible_url_headers = ["google_result_url", "google result url", "google url"]
        
        for h in possible_title_headers:
            if h in header_map:
                google_title_idx = header_map[h]
                logger.info(f"Found Google title column: '{h}' at index {google_title_idx}")
                break
                
        for h in possible_url_headers:
            if h in header_map:
                google_url_idx = header_map[h]
                logger.info(f"Found Google URL column: '{h}' at index {google_url_idx}")
                break
        
        if google_title_idx == -1 and google_url_idx == -1:
             # If strictly looking for google columns and not found, return empty
             logger.warning(f"Google result columns not found. Available headers: {reader.fieldnames}")
             return {}

        # We need original fieldnames to access row by key if using DictReader, 
        # but since we mapped indices, we can iterate rows and use keys from fieldnames
        # actually DictReader rows are dicts. We need the actual keys.
        
        original_headers = reader.fieldnames
        title_key = original_headers[google_title_idx] if google_title_idx != -1 else None
        url_key = original_headers[google_url_idx] if google_url_idx != -1 else None
        
        # We also need query key
        query_key = None
        for h in ["query", "search query"]:
            if h in header_map:
                query_key = original_headers[header_map[h]]
                logger.info(f"Found query column: '{h}' -> key: '{query_key}'")
                break
        
        if not query_key:
             logger.warning(f"Query column not found. Available headers: {reader.fieldnames}")
             return {}

        row_count = 0
        for row in reader:
            row_count += 1
            query = (row.get(query_key) or "").strip()
            google_title = (row.get(title_key) or "").strip() if title_key else ""
            google_url = (row.get(url_key) or "").strip() if url_key else ""
            
            # If query is present, start a new query group
            if query:
                current_query = query
                if current_query not in google_results_map:
                    google_results_map[current_query] = []
            
            # Add result if both title and URL are present (or at least one)
            if current_query and (google_title or google_url):
                google_results_map[current_query].append({
                    "title": google_title,
                    "url": google_url,
                    "rank": len(google_results_map[current_query]) + 1
                })
        
        logger.info(f"Parsed {row_count} rows, found {len(google_results_map)} unique queries with Google results")

    except Exception as e:
        # If parsing fails, log the error and return empty map
        logger.error(f"Error parsing Google results CSV: {e}", exc_info=True)
        return {}
    
    return google_results_map


def parse_google_results_csv(filepath: str) -> Dict[str, List[Dict]]:
    """
    Parse google_results.csv file and return a dictionary mapping queries to their Google results.
    """
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        return parse_google_results_from_content(content)
    except FileNotFoundError:
        raise ValueError(f"Google results CSV file not found: {filepath}")
    except Exception as e:
        raise ValueError(f"Error parsing Google results CSV: {str(e)}")

